"""Turn an uploaded file into LangChain Documents.

PDF             -> one Document per page (splitter joins them and tracks page numbers).
TXT / MD / DOCX -> a single Document, page=None.
"""
from pathlib import Path

from langchain_core.documents import Document


class UnsupportedFileType(ValueError):
    pass


def _load_pdf(path: Path) -> list[Document]:
    import pdfplumber

    docs: list[Document] = []
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                docs.append(Document(page_content=text, metadata={"page": page_no}))
    return docs


def _load_docx(path: Path) -> list[Document]:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts).strip()
    return [Document(page_content=text, metadata={"page": None})] if text else []


def _load_text(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return [Document(page_content=text, metadata={"page": None})] if text else []


LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_text,
    ".md": _load_text,
}


def load_file(path: str | Path) -> list[Document]:
    path = Path(path)
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise UnsupportedFileType(f"Unsupported file type: {path.suffix}")
    docs = loader(path)
    if not docs:
        raise ValueError("No extractable text found (scanned PDFs need OCR first).")
    return docs